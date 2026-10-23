"""Genera docs/Taller_Ranking_Terminos.docx a partir de resultados/*.csv.

Ejecutar despues de `python -m src.ranking`:
    python -m docs.generar_documento
"""
from __future__ import annotations

import csv
import math
import os
import platform

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RES = os.path.join(RAIZ, "resultados")
SALIDA = os.path.join(RAIZ, "docs", "Taller_Ranking_Terminos.docx")


def leer_csv(nombre):
    with open(os.path.join(RES, nombre), encoding="utf-8") as f:
        return list(csv.reader(f))


def tabla(doc, filas, anchos=None, negrita_encabezado=True):
    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(filas):
        for j, valor in enumerate(fila):
            celda = t.cell(i, j)
            celda.text = str(valor)
            for p in celda.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.bold = negrita_encabezado and i == 0
    if anchos:
        for fila in t.rows:
            for j, ancho in enumerate(anchos):
                fila.cells[j].width = Inches(ancho)
    return t


def leyenda(doc, texto):
    p = doc.add_paragraph(texto)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def formula(doc, texto):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto)
    r.font.name = "Cambria Math"
    r.font.size = Pt(11)


def codigo(doc, texto):
    for linea in texto.strip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.3)
        r = p.add_run(linea)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
    doc.add_paragraph()


def main():
    frecuencias = leer_csv("frecuencias.csv")[1:]
    top = leer_csv("top15.csv")
    tiempos = leer_csv("tiempos.csv")
    escala = leer_csv("escalabilidad.csv")

    n = len(frecuencias)
    valores = [int(f) for _, f in frecuencias]
    k = max(valores)
    total = sum(valores)
    unicos_1 = sum(1 for v in valores if v == 1)
    distintas = len(set(valores))
    t = {fila[0]: fila for fila in tiempos[1:]}
    med = {a: float(t[a][5]) for a in t}
    f_corte = int(top[-1][2])
    empatados_corte = [term for term, f in frecuencias if int(f) == f_corte]
    dentro = [fila[1] for fila in top[1:] if int(fila[2]) == f_corte]
    fuera = [e for e in empatados_corte if e not in dentro]

    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)

    # Portada
    for texto, tam, neg in [
        ("Universidad del Quindío", 14, True),
        ("Programa de Ingeniería de Sistemas y Computación", 12, False),
        ("Análisis de Algoritmos — 2026-2", 12, False),
        ("", 12, False),
        ("Seguimiento: Ranking de términos mediante algoritmos de ordenamiento no comparativos", 16, True),
        ("Counting Sort · Radix Sort · Bucket Sort", 12, False),
        ("", 12, False),
        ("[Nombre del estudiante]", 12, False),
        ("[Fecha de entrega]", 12, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(texto)
        r.font.size = Pt(tam)
        r.bold = neg
    doc.add_page_break()

    # 1. Introduccion
    doc.add_heading("1. Introducción", level=1)
    doc.add_paragraph(
        "Este seguimiento es la primera parte del proyecto final del curso, que analiza un corpus de 20 artículos "
        "científicos sobre inteligencia artificial en educación. El objetivo es generar un listado de los términos "
        "más frecuentes de los abstracts del corpus, ordenado con tres algoritmos que no se basan en comparaciones "
        "entre elementos: Counting Sort, Radix Sort y Bucket Sort. Para cada uno se presenta su fundamentación "
        "teórica, sus tiempos de ejecución medidos y el top 15 de términos que produce."
    )

    # 2. Datos y preprocesamiento
    doc.add_heading("2. Fuente de datos y preprocesamiento", level=1)
    doc.add_paragraph(
        "Se procesaron los 20 PDF entregados por el docente. Solo se usó el abstract de cada artículo, que se extrajo "
        "automáticamente (módulo src/corpus/extraccion.py) con la librería pypdf. El abstract se ubica entre la "
        "marca «Abstract» y la primera de las marcas «Keywords» o «1. Introduction». Dos artículos tienen una "
        "maquetación distinta y necesitaron un ajuste específico: el 7 (las palabras clave aparecen antes del "
        "abstract) y el 9 (no tiene palabras clave ni sección numerada). Los abstracts se guardan en "
        "data/abstracts.json para reutilizarlos en los demás requerimientos."
    )
    doc.add_paragraph("El preprocesamiento (src/corpus/preprocesamiento.py) aplica los pasos siguientes:")
    for paso in [
        "Normalización: se unen las ligaduras tipográficas (ﬁ, ﬀ) y las palabras partidas por guion al final de línea.",
        "Tokenización: se pasa todo a minúsculas y se toman secuencias de letras [a-z]+; los números y signos se descartan.",
        "Filtrado: se eliminan las stopwords del inglés y los tokens de menos de 2 letras. El umbral de 2 conserva «ai», "
        "el término central del dominio.",
        "Conteo: se usa una tabla hash término → frecuencia, con costo O(T) esperado para T tokens. La lista resultante "
        "conserva el orden de primera aparición, que es la entrada de los tres algoritmos.",
    ]:
        doc.add_paragraph(paso, style="List Bullet")
    tabla(doc, [
        ["Indicador", "Valor"],
        ["Abstracts procesados", "20"],
        ["Tokens útiles (después del filtrado)", str(total)],
        ["Términos distintos (n)", str(n)],
        ["Frecuencia máxima (k)", str(k)],
        ["Valores de frecuencia distintos", str(distintas)],
        ["Términos con frecuencia 1", f"{unicos_1} ({100 * unicos_1 / n:.1f} %)"],
    ], anchos=[3.2, 1.6])
    leyenda(doc, "Tabla 1. Características de la entrada que reciben los algoritmos.")
    doc.add_paragraph(
        f"La distribución de frecuencias es muy sesgada, como predice la ley de Zipf: el {100 * unicos_1 / n:.0f} % de los "
        f"términos aparece una sola vez y solo unos pocos superan las 30 apariciones. Este dato es importante para el "
        f"análisis porque k = {k} es mucho menor que n = {n}, y en esa situación los algoritmos no comparativos se "
        f"comportan como lineales."
    )

    # 3. Fundamentacion teorica
    doc.add_heading("3. Fundamentación teórica", level=1)
    doc.add_heading("3.1 La cota Ω(n log n) y cómo se evita", level=2)
    doc.add_paragraph(
        "Cualquier algoritmo de ordenamiento basado en comparaciones (Merge Sort, Quick Sort, Heap Sort) puede "
        "modelarse como un árbol de decisión binario. Cada nodo interno es una comparación ai ≤ aj y cada hoja es una "
        "de las n! permutaciones posibles de la entrada. Un árbol binario de altura h tiene como máximo 2^h hojas, "
        "por lo tanto:"
    )
    formula(doc, "2^h ≥ n!   ⇒   h ≥ log₂(n!) = Θ(n log n)   (aproximación de Stirling)")
    doc.add_paragraph(
        "Así, en el peor caso todo ordenamiento por comparación hace Ω(n log n) comparaciones. Los tres algoritmos de "
        "este trabajo no están sujetos a esa cota porque nunca comparan dos elementos entre sí: usan el valor de la "
        "clave (la frecuencia, un entero acotado) como índice de un arreglo. Esa operación de indexación en O(1) "
        "aporta más información que una comparación binaria, y por eso se puede bajar a tiempo lineal. A cambio, el "
        "algoritmo depende de que las claves sean enteros (o se puedan transformar en enteros) dentro de un rango "
        "conocido."
    )
    doc.add_paragraph(
        "Para que los tres algoritmos produzcan exactamente el mismo ranking, se implementaron como estables y "
        "descendentes. Estable significa que los términos con la misma frecuencia conservan su orden de primera "
        "aparición en el corpus."
    )

    doc.add_heading("3.2 Counting Sort — O(n + k)", level=2)
    doc.add_paragraph("El algoritmo tiene cuatro fases (src/ordenamiento/counting_sort.py):")
    for paso in [
        "Hallar la clave máxima k: un recorrido de la entrada, O(n).",
        "Construir el histograma conteo[f]: un recorrido de la entrada, O(n), sobre un arreglo de tamaño k + 1, O(k).",
        "Calcular la posición inicial de cada clave con una suma acumulada desde k hasta 0 (de mayor a menor por ser "
        "descendente): O(k).",
        "Colocar cada elemento en salida[inicio[f]] e incrementar inicio[f]: O(n). Como se recorre la entrada de "
        "izquierda a derecha, los empates conservan su orden y el algoritmo es estable.",
    ]:
        doc.add_paragraph(paso, style="List Number")
    formula(doc, "T(n, k) = c₁·n + c₂·k + c₃·n + c₄·k = O(n + k)          S(n, k) = O(n + k)")
    doc.add_paragraph(
        f"Si k = O(n), el algoritmo es O(n). En este corpus k = {k} y n = {n}, así que el término k pesa poco y el "
        f"tiempo depende casi solo de n. El algoritmo dejaría de ser conveniente si k fuera mucho mayor que n (por "
        f"ejemplo, si las claves fueran identificadores de 10⁹), porque el arreglo de conteo dominaría el costo."
    )
    codigo(doc, """
for clave in range(k, -1, -1):        # suma acumulada descendente, O(k)
    inicio[clave] = acumulado
    acumulado += conteo[clave]
for par in pares:                     # colocacion estable, O(n)
    salida[inicio[par[1]]] = par
    inicio[par[1]] += 1
""")

    doc.add_heading("3.3 Radix Sort (LSD) — O(d·(n + b)) = O(n)", level=2)
    d = len(str(k))
    doc.add_paragraph(
        "Radix Sort ordena la clave dígito por dígito en base b, empezando por el menos significativo (LSD). Cada "
        "pasada es un Counting Sort estable sobre un solo dígito, cuyas claves están en el rango [0, b−1], así que "
        "cuesta O(n + b). Se necesitan d = ⌊log_b(k)⌋ + 1 pasadas:"
    )
    formula(doc, "T(n) = d · O(n + b) = O(d·(n + b))")
    doc.add_paragraph(
        f"Con b = 10 y k = {k}, se obtiene d = {d} pasadas. Como d y b son constantes que no crecen con el número de "
        f"términos, el costo es O(n). El algoritmo es correcto porque cada pasada es estable: después de ordenar por "
        f"las unidades, ordenar por las decenas deja juntas las claves con la misma decena sin alterar el orden de las "
        f"unidades. Por inducción sobre el número de pasadas, al terminar la lista queda ordenada por la clave "
        f"completa. Frente a Counting Sort, Radix Sort no necesita un arreglo de tamaño k, solo uno de tamaño b. "
        f"Por eso conviene cuando k es enorme, aunque aquí, con k pequeño, paga d recorridos completos en lugar de uno."
    )

    doc.add_heading("3.4 Bucket Sort — O(n + m) en promedio", level=2)
    doc.add_paragraph(
        "Bucket Sort reparte los n elementos en m cubetas según el intervalo de su clave. La cubeta i recibe las "
        "claves con ⌊(f − mín)·m / (máx − mín + 1)⌋ = i. Después ordena cada cubeta con Insertion Sort (que es "
        "estable) y concatena las cubetas de la mayor a la menor. Si nᵢ es el número de elementos de la cubeta i:"
    )
    formula(doc, "T(n) = O(n) [distribución] + O(m) [recorrido de cubetas] + Σ O(nᵢ²) [Insertion Sort]")
    doc.add_paragraph(
        "Si las claves se distribuyen de manera uniforme y m = Θ(n), cada cubeta recibe en promedio O(1) elementos y "
        "E[Σ nᵢ²] = O(n), así que el costo esperado es O(n). En el peor caso todas las claves caen en una sola "
        "cubeta y el costo sube a O(n²)."
    )
    doc.add_paragraph(
        f"Las frecuencias de este corpus no son uniformes (siguen la ley de Zipf), así que la hipótesis del caso "
        f"promedio no se cumple. Aun así, el algoritmo es lineal en este caso por otra razón: se usan m = n = {n} "
        f"cubetas para un rango de solo {k} valores, de modo que cada cubeta cubre menos de un valor de frecuencia y "
        f"todos los elementos de una cubeta tienen la misma clave. El Insertion Sort hace entonces una sola "
        f"comparación por elemento y ningún desplazamiento, por ejemplo en la cubeta de frecuencia 1, que tiene "
        f"{unicos_1} términos. Con pocas cubetas (por ejemplo m = 5), las de menor rango mezclarían frecuencias "
        f"distintas y el Insertion Sort sí haría desplazamientos."
    )

    tabla(doc, [
        ["Algoritmo", "Tiempo", "Espacio extra", "Estable", "Requisito sobre las claves"],
        ["Counting Sort", "O(n + k)", "O(n + k)", "Sí", "Enteros en [0, k] con k = O(n)"],
        ["Radix Sort", "O(d·(n + b))", "O(n + b)", "Sí", "Enteros o cadenas de d dígitos"],
        ["Bucket Sort", "O(n + m) prom. / O(n²) peor", "O(n + m)", "Sí*", "Distribución aprox. uniforme"],
    ], anchos=[1.1, 1.6, 1.0, 0.6, 2.2])
    leyenda(doc, "Tabla 2. Resumen de complejidades. *Estable si el ordenamiento interno es estable (Insertion Sort lo es).")

    # 4. Medicion empirica
    doc.add_heading("4. Medición empírica", level=1)
    doc.add_paragraph(
        f"Cada algoritmo se ejecutó {tiempos[1][3]} veces sobre la lista completa de frecuencias del corpus "
        f"(n = {n}, k = {k}), después de una ejecución de calentamiento. El tiempo se midió con "
        f"time.perf_counter_ns() y se reporta en milisegundos. Se presenta la mediana porque es robusta frente a las "
        f"interrupciones del sistema operativo, que inflan algunas ejecuciones aisladas (ver la columna «máx»)."
    )
    filas = [["Algoritmo", "mín (ms)", "mediana (ms)", "media (ms)", "desv. (ms)", "máx (ms)"]]
    for fila in tiempos[1:]:
        filas.append([fila[0], fila[4], fila[5], fila[6], fila[7], fila[8]])
    tabla(doc, filas, anchos=[1.3, 0.9, 1.0, 0.9, 0.9, 0.9])
    leyenda(doc, f"Tabla 3. Tiempos de ejecución sobre el corpus completo ({tiempos[1][3]} repeticiones).")
    doc.add_picture(os.path.join(RES, "tiempos.png"), width=Inches(6))
    leyenda(doc, "Figura 1. Mediana del tiempo por algoritmo; las barras de error marcan el rango intercuartílico.")

    rapido = min(med, key=med.get)
    doc.add_paragraph(
        f"Para comprobar empíricamente que el crecimiento es lineal, se repitió la medición sobre entradas sintéticas "
        f"de tamaño 1n a 32n, que se construyen replicando la lista del corpus. Así se conserva la distribución de "
        f"frecuencias y el valor de k, y solo crece n."
    )
    filas = [escala[0]] + escala[1:]
    tabla(doc, filas, anchos=[0.7, 0.9, 1.4, 1.4, 1.4])
    leyenda(doc, "Tabla 4. Mediana del tiempo (ms) al aumentar el tamaño de la entrada.")
    doc.add_picture(os.path.join(RES, "escalabilidad.png"), width=Inches(6))
    leyenda(doc, "Figura 2. Tiempo frente a n: las tres curvas son aproximadamente rectas.")

    ratios = []
    for j, a in enumerate(["Counting Sort", "Radix Sort", "Bucket Sort"]):
        ratios.append(f"{a} ×{float(escala[-1][2 + j]) / float(escala[1][2 + j]):.1f}")
    doc.add_paragraph(
        "Análisis: al multiplicar n por 32, los tiempos se multiplican por aproximadamente "
        + ", ".join(ratios)
        + ". Estos factores están cerca del ×32 que predice un costo Θ(n) y por debajo del "
        f"×{32 * math.log(32 * n) / math.log(n):.0f} que correspondería a Θ(n log n) en este rango "
        "(32·log(32n)/log n). La diferencia respecto a 32 se explica por factores que el modelo asintótico no "
        "considera: con entradas grandes las listas ya no caben en la caché del procesador, y el recolector de "
        "memoria de Python interviene con más frecuencia. Bucket Sort es el más afectado, porque crea una lista "
        f"por cubeta. En cuanto a las constantes, {rapido} es el más rápido porque hace un solo recorrido de "
        f"distribución con un arreglo auxiliar pequeño (k + 1 = {k + 1} posiciones). Radix Sort hace {d} pasadas "
        f"completas, cada una con su propio arreglo de salida, lo que explica que tarde más o menos el doble. Bucket "
        f"Sort crea {n} listas de Python (una por cubeta) y las concatena, y ese costo de gestión de memoria es "
        f"mayor que el de un arreglo de enteros. Los tres pertenecen a la misma clase asintótica, y la diferencia "
        f"entre ellos está en el factor constante."
    )

    # 5. Validacion
    doc.add_heading("5. Validación de resultados: top 15", level=1)
    doc.add_paragraph(
        "La tabla 5 muestra los 15 términos más frecuentes que obtuvo cada algoritmo. Además, el programa compara "
        "la salida completa de cada algoritmo con la de la función sorted() de Python (usada solo como referencia de "
        "verificación). Las tres coinciden elemento por elemento, incluido el orden de los empates."
    )
    filas = [["#", "Counting Sort", "f", "Radix Sort", "f", "Bucket Sort", "f"]] + top[1:]
    tabla(doc, filas, anchos=[0.35, 1.3, 0.45, 1.3, 0.45, 1.3, 0.45])
    leyenda(doc, "Tabla 5. Top 15 de términos obtenido con cada algoritmo (f = frecuencia).")
    doc.add_picture(os.path.join(RES, "top15.png"), width=Inches(6))
    leyenda(doc, "Figura 3. Top 15 de términos del corpus (idéntico para los tres algoritmos).")
    if fuera:
        doc.add_paragraph(
            f"Empate en el corte: hay {len(empatados_corte)} términos con frecuencia {f_corte} "
            f"({', '.join(empatados_corte)}). Como los algoritmos son estables, entran al top 15 los que aparecen "
            f"primero en el corpus ({', '.join(dentro)}), y «{', '.join(fuera)}» queda en la posición 16. Si los "
            f"algoritmos no fueran estables, el contenido del top 15 podría variar entre ejecuciones o entre "
            f"algoritmos."
        )
    doc.add_paragraph(
        "Interpretación: los términos dominantes (ai, education, research, learning, aied, review, assessment) "
        "describen bien el corpus, que está formado por revisiones de literatura sobre IA en educación y "
        "evaluación. Los términos «artificial» e «intelligence» aparecen por separado porque el análisis es por "
        "palabras individuales. Contar bigramas, o unificar variantes como education/educational o study/studies "
        "mediante lematización, son mejoras posibles que no cambian el análisis de los algoritmos."
    )

    # 6. Conclusiones
    doc.add_heading("6. Conclusiones", level=1)
    for c in [
        "Los tres algoritmos superan la cota Ω(n log n) porque usan la clave como índice en lugar de comparar "
        "elementos. Para eso necesitan claves enteras y acotadas, condición que las frecuencias cumplen.",
        f"Las mediciones confirman el crecimiento lineal: los tiempos escalan de forma proporcional a n. En el "
        f"corpus real los tres algoritmos ordenan {n} términos en menos de 1 ms.",
        f"Counting Sort es el más eficiente para este problema porque k = {k} ≪ n. Radix Sort sería preferible si "
        f"las claves tuvieran un rango muy grande, y Bucket Sort depende de la distribución de las claves y de la "
        f"elección de m.",
        "La estabilidad es necesaria para que el ranking sea reproducible cuando hay empates de frecuencia.",
    ]:
        doc.add_paragraph(c, style="List Bullet")

    # 7. Reproducibilidad y uso de IA
    doc.add_heading("7. Reproducibilidad", level=1)
    doc.add_paragraph(
        f"Entorno de medición: Python {platform.python_version()}, {platform.system()} {platform.release()}, "
        f"procesador {platform.processor() or platform.machine()}. Comandos:"
    )
    codigo(doc, """
pip install -r requirements.txt
python -m src.corpus.extraccion          # extrae los abstracts a data/abstracts.json
python -m src.ranking                    # ordena, mide y genera resultados/
python -m unittest discover -s tests     # pruebas de los algoritmos
""")
    doc.add_heading("8. Declaración de uso de IA generativa", level=1)
    doc.add_paragraph(
        "[Completar según los lineamientos del docente. Por ejemplo: se usó Claude (Anthropic) como asistente para "
        "estructurar el código y redactar el borrador de este documento. Los algoritmos se implementaron con "
        "estructuras básicas del lenguaje, sin funciones de ordenamiento de alto nivel, y el estudiante revisó y "
        "validó el análisis.]"
    )

    doc.save(SALIDA)
    print(f"Documento generado: {SALIDA}")


if __name__ == "__main__":
    main()
